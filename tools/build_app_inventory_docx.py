from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "DAKETI_APP_INVENTORY.md"
OUTPUT = ROOT / "deliverables" / "Daketi_Phase_I_Full_App_Inventory.docx"
ICON = ROOT / "assets" / "images" / "app_icon.png"

BLACK = "000000"
DARK = "262626"
MID = "666666"
LIGHT = "F2F2F2"
PALE = "FAFAFA"
ORANGE = "E97822"
BORDER = "D9D9D9"


def set_font(run, name="Arial", size=None, bold=None, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:color"), BORDER)


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:keepNext")) is None:
        p_pr.append(OxmlElement("w:keepNext"))


def keep_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_inline(paragraph, text, base_size=10.5, base_color=BLACK):
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_font(run, size=base_size, color=base_color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, size=base_size, bold=True, color=base_color)
        else:
            run = paragraph.add_run(token[1:-1])
            set_font(run, name="Courier New", size=9, color=DARK)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_font(run, size=base_size, color=base_color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, size=8.5, color=MID)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def set_repeat_table_headers(table):
    if table.rows:
        repeat_header(table.rows[0])


def add_table(doc, rows):
    headers = rows[0]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    hdr = table.rows[0]
    repeat_header(hdr)
    keep_row(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        shade_cell(cell, DARK)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        set_font(run, size=9, bold=True, color="FFFFFF")
    for ridx, values in enumerate(rows[1:]):
        cells = table.add_row().cells
        keep_row(table.rows[-1])
        for idx, text in enumerate(values):
            cell = cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if ridx % 2:
                shade_cell(cell, PALE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, text, base_size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size, before, after in (
        ("Title", 27, 0, 12),
        ("Heading 1", 17, 16, 7),
        ("Heading 2", 13, 12, 5),
        ("Heading 3", 11, 9, 4),
    ):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        # Remove any built-in theme border, especially Word's Title blue rule.
        p_pr = style.element.get_or_add_pPr()
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is not None:
            p_pr.remove(p_bdr)


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    # Cover page
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(38)
    if ICON.exists():
        p.add_run().add_picture(str(ICON), width=Inches(1.1))
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Daketi Phase I Full App Inventory")
    set_font(run, size=27, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("Screens  Controls  Text Fields  Sounds  Images  Implementation Status")
    set_font(run, size=11.5, color=MID)
    intro = doc.add_paragraph()
    intro.paragraph_format.left_indent = Inches(0.6)
    intro.paragraph_format.right_indent = Inches(0.6)
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(
        intro,
        "This document provides a complete inventory of the current Daketi Phase I Flutter application. It records the visible app structure, interface controls, input fields, packaged audio and image assets, and known placeholders.",
        base_size=11,
    )
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(25)
    add_inline(meta, "Audit date  16 September 2026", base_size=9.5, base_color=MID)
    meta2 = doc.add_paragraph()
    meta2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(meta2, "Landscape Flutter application", base_size=9.5, base_color=MID)
    doc.add_page_break()

    # A quiet page-number footer keeps the report easy to reference.
    add_page_number(section.footer.paragraphs[0])

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    # Skip source title and metadata; the cover replaces them.
    start = next(i for i, line in enumerate(lines) if line.startswith("## 1."))
    lines = lines[start:]
    i = 0
    while i < len(lines):
        raw = lines[i].rstrip()
        text = raw.strip()
        if not text or text == "---":
            i += 1
            continue
        if text.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                parts = [p.strip() for p in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", p) for p in parts):
                    rows.append(parts)
                i += 1
            add_table(doc, rows)
            continue
        if text.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, text[3:], base_size=17)
            keep_with_next(p)
        elif text.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, text[4:], base_size=13)
            keep_with_next(p)
        elif text.startswith("#### "):
            p = doc.add_paragraph(style="Heading 3")
            add_inline(p, text[5:], base_size=11)
            keep_with_next(p)
        elif re.match(r"^\d+\.\s+", text):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\.\s+", "", text))
            p.paragraph_format.space_after = Pt(3)
        elif text.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, text[2:])
            p.paragraph_format.space_after = Pt(3)
        else:
            p = doc.add_paragraph()
            add_inline(p, text)
        i += 1

    # Set document metadata without personal author data.
    props = doc.core_properties
    props.title = "Daketi Phase I Full App Inventory"
    props.subject = "Application screens controls assets and implementation status"
    props.author = "Daketi"
    props.keywords = "Daketi, Flutter, app inventory, UI, sounds, images"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
